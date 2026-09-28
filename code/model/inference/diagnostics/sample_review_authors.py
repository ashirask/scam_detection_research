#!/usr/bin/env python3
"""Create a reproducible Excel review workbook from flagged authors."""

import argparse
import hashlib
import logging
import random
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import orjson
from openpyxl.styles import PatternFill


LOGGER = logging.getLogger(__name__)
ANNOTATION_COLUMNS = [
    "reviewer_1_label",
    "reviewer_1_notes",
    "reviewer_2_label",
    "reviewer_2_notes",
    "adjudicated_label",
]
COMMENT_COLUMNS = [
    "author",
    "comment_index",
    "comment_id",
    "body",
    "subreddit",
    "created_utc",
    "score",
]
SUBMISSION_COLUMNS = [
    "author",
    "submission_index",
    "submission_id",
    "title",
    "selftext",
    "subreddit",
    "created_utc",
    "score",
]
AUTHOR_COLORS = [
    "F4CCCC", "D9EAD3", "CFE2F3", "FFF2CC", "D9D2E9",
    "FCE5CD", "D0E0E3", "EAD1DC", "C9DAF8", "B6D7A8",
    "F9CB9C", "CFE2F3", "D5A6BD", "A2C4C9", "FFE599",
    "B4A7D6", "F4B183", "A9D18E", "9DC3E6", "FFD966",
]


def setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_flagged_authors(path: Path) -> pd.DataFrame:
    if not path.is_file():
        raise FileNotFoundError(f"Flagged-author CSV not found: {path}")

    flagged = pd.read_csv(path)
    if "author" not in flagged.columns:
        raise ValueError("Flagged-author CSV must contain an 'author' column")

    authors = flagged["author"].astype("string")
    if authors.isna().any() or authors.str.strip().eq("").any():
        raise ValueError("Flagged-author CSV contains missing or empty author values")
    if authors.duplicated().any():
        duplicates = sorted(authors[authors.duplicated()].unique().tolist())
        raise ValueError(f"Flagged-author CSV contains duplicate authors: {duplicates[:10]}")

    flagged["author"] = authors
    return flagged


def filter_authors_by_total_posts(
    flagged: pd.DataFrame,
    min_total_posts: int
) -> tuple[pd.DataFrame, int]:
    """
    Filter out authors with fewer than min_total_posts total posts.

    Uses the 'total_posts' column from flagged_authors.csv (which includes
    both comments and submissions) to filter efficiently without loading
    the raw bundle.

    Args:
        flagged: DataFrame of flagged authors (includes total_posts column)
        min_total_posts: Minimum total posts required

    Returns:
        tuple: (filtered_flagged, n_filtered)
            filtered_flagged: DataFrame with authors meeting threshold
            n_filtered: Number of authors filtered out
    """
    if "total_posts" not in flagged.columns:
        raise ValueError("Flagged-author CSV must contain a 'total_posts' column for filtering")

    n_before = len(flagged)
    filtered = flagged[flagged["total_posts"] >= min_total_posts].copy()
    n_filtered = n_before - len(filtered)

    if len(filtered) == 0:
        raise ValueError(
            f"All authors were filtered out. Minimum total posts ({min_total_posts}) "
            f"is higher than all available authors. Maximum total posts: {flagged['total_posts'].max()}"
        )

    return filtered, n_filtered


def select_authors(flagged: pd.DataFrame, n: int, seed: int, allow_fewer: bool) -> pd.DataFrame:
    if n <= 0:
        raise ValueError("--n must be greater than zero")
    available = len(flagged)
    if n > available and not allow_fewer:
        raise ValueError(
            f"Requested {n} authors but only {available} are available; "
            "use --allow-fewer to accept a smaller sample"
        )

    sample_size = min(n, available)
    author_order = sorted(flagged["author"].tolist())
    rng = random.Random(seed)
    selected_authors = set(rng.sample(author_order, sample_size))
    selected = flagged[flagged["author"].isin(selected_authors)].copy()
    selected = selected.sort_values("author", kind="stable").reset_index(drop=True)
    return selected


def load_raw_bundle(path: Path | None, selected_authors: set[str]):
    records = {}
    if path is None:
        return records
    if not path.is_file():
        raise FileNotFoundError(f"Raw bundle not found: {path}")

    with path.open("rb") as source:
        for line_number, line in enumerate(source, start=1):
            if not line.strip():
                continue
            try:
                record = orjson.loads(line)
            except orjson.JSONDecodeError as exc:
                raise ValueError(f"Invalid JSON on line {line_number} of {path}: {exc}") from exc
            author = record.get("author")
            if author in selected_authors:
                if author in records:
                    raise ValueError(f"Raw bundle contains duplicate records for author '{author}'")
                records[author] = record
    return records


def flatten_comments(records: dict[str, dict]) -> pd.DataFrame:
    rows = []
    for author in sorted(records):
        for index, comment in enumerate(records[author].get("comments", []), start=1):
            rows.append(
                {
                    "author": author,
                    "comment_index": index,
                    "comment_id": comment.get("id"),
                    "body": comment.get("body", ""),
                    "subreddit": comment.get("subreddit", ""),
                    "created_utc": comment.get("created_utc"),
                    "score": comment.get("score"),
                }
            )
    return pd.DataFrame(rows, columns=COMMENT_COLUMNS)


def flatten_submissions(records: dict[str, dict]) -> pd.DataFrame:
    rows = []
    for author in sorted(records):
        for index, submission in enumerate(records[author].get("submissions", []), start=1):
            rows.append(
                {
                    "author": author,
                    "submission_index": index,
                    "submission_id": submission.get("id"),
                    "title": submission.get("title", ""),
                    "selftext": submission.get("selftext", ""),
                    "subreddit": submission.get("subreddit", ""),
                    "created_utc": submission.get("created_utc"),
                    "score": submission.get("score"),
                }
            )
    return pd.DataFrame(rows, columns=SUBMISSION_COLUMNS)


def write_filtered_bundle(path: Path, records: dict[str, dict], selected_authors: list[str]) -> None:
    with path.open("wb") as output:
        for author in selected_authors:
            record = records.get(author)
            if record is not None:
                output.write(orjson.dumps(record) + b"\n")


def infer_threshold(flagged_path: Path) -> str:
    match = re.search(r"threshold_([0-9.]+)", str(flagged_path.parent))
    return match.group(1) if match else "unknown"


def build_metadata(
    flagged_path: Path,
    raw_bundle_path: Path | None,
    output_path: Path,
    available_count: int,
    selected: pd.DataFrame,
    records: dict[str, dict],
    n: int,
    seed: int,
    allow_fewer: bool,
    author_colors: dict[str, str],
    min_total_posts: int | None = None,
    n_flagged_before_filter: int | None = None,
    n_flagged_after_filter: int | None = None,
    n_authors_filtered: int | None = None,
) -> pd.DataFrame:
    selected_authors = selected["author"].tolist()
    missing = sorted(set(selected_authors) - set(records)) if raw_bundle_path else selected_authors
    model = selected["model"].iloc[0] if "model" in selected.columns and len(selected) else "unknown"
    strategy = selected["strategy"].iloc[0] if "strategy" in selected.columns and len(selected) else "unknown"
    scored_at = selected["scored_at"].iloc[0] if "scored_at" in selected.columns and len(selected) else "unknown"

    rows = [
        ("generated_at_utc", datetime.now(timezone.utc).isoformat()),
        ("flagged_authors_path", str(flagged_path.resolve())),
        ("flagged_authors_sha256", sha256_file(flagged_path)),
        ("raw_bundle_path", str(raw_bundle_path.resolve()) if raw_bundle_path else "not provided"),
        ("output_workbook", str(output_path.resolve())),
        ("model", str(model)),
        ("strategy", str(strategy)),
        ("threshold", infer_threshold(flagged_path)),
        ("scored_at", str(scored_at)),
        ("requested_sample_size", n),
        ("available_author_count", available_count),
        ("selected_author_count", len(selected)),
        ("random_seed", seed),
        ("allow_fewer", allow_fewer),
        ("raw_records_found", len(records)),
        ("raw_records_missing", len(missing)),
        ("missing_raw_authors", ", ".join(missing)),
        ("review_label_values", "bot, human, uncertain"),
        ("review_scope", "Flagged-author review supports false-positive investigation; it is not ground truth."),
        ("raw_bundle_scope", "The existing bundle contains at most 20 comments and 20 submissions per author."),
        ("selected_authors", ", ".join(selected_authors)),
    ]

    if min_total_posts is not None:
        rows.extend([
            ("min_total_posts_threshold", str(min_total_posts)),
            ("n_flagged_before_filter", str(n_flagged_before_filter)),
            ("n_flagged_after_filter", str(n_flagged_after_filter)),
            ("n_authors_filtered", str(n_authors_filtered)),
        ])

    rows.extend(
        (f"author_color_{index}", f"{author} = #{author_colors[author]}")
        for index, author in enumerate(selected_authors, start=1)
    )
    return pd.DataFrame(rows, columns=["field", "value"])


def build_author_colors(authors: list[str]) -> dict[str, str]:
    return {
        author: AUTHOR_COLORS[index % len(AUTHOR_COLORS)]
        for index, author in enumerate(sorted(authors))
    }


def write_workbook(
    output_path: Path,
    selected: pd.DataFrame,
    comments: pd.DataFrame,
    submissions: pd.DataFrame,
    metadata: pd.DataFrame,
    author_colors: dict[str, str],
) -> None:
    authors = selected.copy()
    for column in ANNOTATION_COLUMNS:
        authors[column] = ""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        authors.to_excel(writer, sheet_name="Authors_Features", index=False)
        comments.to_excel(writer, sheet_name="Comments", index=False)
        submissions.to_excel(writer, sheet_name="Submissions", index=False)
        metadata.to_excel(writer, sheet_name="Metadata", index=False)

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for column_cells in worksheet.columns:
                width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 10), 60)
                worksheet.column_dimensions[column_cells[0].column_letter].width = width

        for sheet_name in ("Comments", "Submissions"):
            worksheet = writer.book[sheet_name]
            author_column = next(
                cell.column for cell in worksheet[1] if cell.value == "author"
            )
            for row in worksheet.iter_rows(min_row=2):
                author = row[author_column - 1].value
                color = author_colors.get(author)
                if color:
                    fill = PatternFill(fill_type="solid", fgColor=color)
                    for cell in row:
                        cell.fill = fill


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Sample flagged authors and create a multi-sheet Excel review workbook."
    )
    parser.add_argument("--flagged-authors", required=True, type=Path)
    parser.add_argument("--raw-bundle", type=Path, help="Optional raw_posts_bundle.jsonl from Phase C2")
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--n", type=int, default=20, help="Number of authors to sample (default: 20)")
    parser.add_argument("--seed", type=int, default=20260908, help="Random seed (default: 20260908)")
    parser.add_argument("--allow-fewer", action="store_true")
    parser.add_argument("--workbook-name", default="sampled_review.xlsx")
    parser.add_argument("--min-total-posts", type=int, required=True,
                        help="Minimum total posts (comments + submissions) required")
    return parser.parse_args()


def main() -> None:
    setup_logging()
    args = parse_args()
    flagged_path = args.flagged_authors.resolve()
    raw_bundle_path = args.raw_bundle.resolve() if args.raw_bundle else None
    output_dir = args.output_dir.resolve()
    workbook_path = output_dir / args.workbook_name
    filtered_bundle_path = output_dir / "sampled_raw_posts_bundle.jsonl"

    flagged = load_flagged_authors(flagged_path)
    flagged_filtered, n_filtered = filter_authors_by_total_posts(flagged, args.min_total_posts)
    selected = select_authors(flagged_filtered, args.n, args.seed, args.allow_fewer)
    selected_authors = selected["author"].tolist()
    author_colors = build_author_colors(selected_authors)
    records = load_raw_bundle(raw_bundle_path, set(selected_authors))
    comments = flatten_comments(records)
    submissions = flatten_submissions(records)
    metadata = build_metadata(
        flagged_path,
        raw_bundle_path,
        workbook_path,
        len(flagged_filtered),
        selected,
        records,
        args.n,
        args.seed,
        args.allow_fewer,
        author_colors,
        min_total_posts=args.min_total_posts,
        n_flagged_before_filter=len(flagged),
        n_flagged_after_filter=len(flagged_filtered),
        n_authors_filtered=n_filtered,
    )

    write_workbook(workbook_path, selected, comments, submissions, metadata, author_colors)
    if raw_bundle_path:
        write_filtered_bundle(filtered_bundle_path, records, selected_authors)

    LOGGER.info("Filtered %d authors from %d flagged authors (min_total_posts: %d)",
                n_filtered, len(flagged), args.min_total_posts)
    LOGGER.info("Selected %d authors from %d filtered authors", len(selected), len(flagged_filtered))
    LOGGER.info("Raw records found for %d selected authors", len(records))
    LOGGER.info("Workbook written to %s", workbook_path)
    if raw_bundle_path:
        LOGGER.info("Filtered raw bundle written to %s", filtered_bundle_path)


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError) as exc:
        LOGGER.error(str(exc))
        sys.exit(1)
